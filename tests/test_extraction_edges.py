"""Failure-path tests for extraction and the store's cleanup branches.

``tests/test_extraction.py`` covers the happy paths and the DOCX hostile cases;
this file covers what it does not: the PDF failure family (unopenable document,
one unreadable page inside a readable one), the ZIP-sniff fallback that runs when
a DOCX's central directory is gone, the decode ladder's last rung, and the two
branches of ``document_store`` that only run when a filesystem call fails.

It also pins Deviation 32 -- that a document reaching the detector arrives
hard-wrapped -- on the two container formats it is claimed for, because a later
"tidy up the extracted text" change to ``_normalise`` would silently falsify a
deviation that three other modules are written against.

Every fixture is synthesised in-module from invented strings; no binaries are
committed and nothing here resembles a real record. ``asyncio_mode = "auto"``
is set in ``pyproject.toml``, so ``async def test_*`` needs no decorator.
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
from collections.abc import Collection, Sequence
from pathlib import Path
from typing import Any

import pytest
from docx import Document as new_docx

from secure_context_pipeline.config import Settings
from secure_context_pipeline.errors import PipelineError
from secure_context_pipeline.store import document_store
from secure_context_pipeline.store.document_store import (
    RAW_SUFFIX,
    TEXT_SUFFIX,
    EncryptedDocumentStore,
)
from secure_context_pipeline.store.extraction import KIND_DOCX, detect_kind, extract_text

EXTRACTION_LOGGER = "secure_context_pipeline.store.extraction"
STORE_LOGGER = "secure_context_pipeline.store.document_store"

#: A structured detail may name an exception *class*, never its message: a pypdf
#: or lxml message can quote raw stream bytes, which at this point in the
#: pipeline are unredacted document content.
_CLASS_NAME_RE = re.compile(r"\A[A-Za-z_][A-Za-z0-9_]*\Z")

#: Invented outright -- no real company, no real person, no real record.
DOC_ID_RE = re.compile(r"\A[0-9a-f]{32}\Z")
SYNTHETIC_PARAGRAPH = "Synthetic Northwind quarterly note, fictional throughout."


# ---------------------------------------------------------------------------
# fixture builders
# ---------------------------------------------------------------------------


def build_pdf_bytes(pages: Sequence[Sequence[str]], *, broken_pages: Collection[int] = ()) -> bytes:
    """A multi-page PDF drawing ``lines`` on each page, one text line per row.

    Written by hand because no dependency here exposes a text-drawing API. Each
    line gets its own ``Tm`` at a lower ``y``, which is exactly how a real
    renderer lays out wrapped prose and is what makes pypdf emit one output line
    per drawn line.

    A page listed in ``broken_pages`` gets ``/Resources 7`` -- a number where the
    spec requires a dictionary. The file still opens and still has the right page
    count; the damage only surfaces when *that page's* text is extracted, which
    is the shape ``_extract_pdf``'s per-page guard exists for.
    """
    page_count = len(pages)
    font_number = 3 + 2 * page_count

    objects: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    kids = b" ".join(f"{3 + 2 * index} 0 R".encode("ascii") for index in range(page_count))
    objects.append(
        b"<< /Type /Pages /Kids [" + kids + b"] /Count " + str(page_count).encode("ascii") + b" >>"
    )

    for index, lines in enumerate(pages):
        drawn: list[bytes] = []
        y = 700
        for line in lines:
            escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
            drawn.append(f"1 0 0 1 72 {y} Tm ({escaped}) Tj".encode("ascii"))
            y -= 16
        stream = b"BT /F1 12 Tf " + b" ".join(drawn) + b" ET"
        resources = (
            b"7"
            if index in broken_pages
            else b"<< /Font << /F1 " + str(font_number).encode("ascii") + b" 0 R >> >>"
        )
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources "
            + resources
            + b" /Contents "
            + str(4 + 2 * index).encode("ascii")
            + b" 0 R >>"
        )
        objects.append(
            b"<< /Length "
            + str(len(stream)).encode("ascii")
            + b" >>\nstream\n"
            + stream
            + b"\nendstream"
        )

    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode("ascii") + body + b"\nendobj\n"
    xref_offset = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode("ascii")
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode("ascii")
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n".encode("ascii")
    out += f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    return bytes(out)


def build_simple_docx_bytes() -> bytes:
    """A one-paragraph .docx, in memory."""
    document = new_docx()
    document.add_paragraph(SYNTHETIC_PARAGRAPH)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def warnings_from(caplog: pytest.LogCaptureFixture, logger_name: str) -> list[str]:
    """Rendered WARNING messages from one logger -- pypdf logs warnings of its own."""
    return [
        record.getMessage()
        for record in caplog.records
        if record.name == logger_name and record.levelno >= logging.WARNING
    ]


# ---------------------------------------------------------------------------
# 1. PDF: the document cannot be opened at all
# ---------------------------------------------------------------------------


async def test_unopenable_pdf_raises_pdf_parse_failed_naming_only_the_class() -> None:
    """A blob with PDF magic but no usable structure fails structurally, not with a traceback.

    ``error`` must be a bare class name: pypdf's messages quote the bytes it
    choked on, and those bytes are unredacted document content at this point in
    the pipeline.
    """
    canary = "SYNTHETIC-CANARY-PDF-4471"
    payload = b"%PDF-1.4\n" + canary.encode("ascii") + b"\x00\x01\x02" * 40

    with pytest.raises(PipelineError) as excinfo:
        await extract_text("shredded.pdf", payload)

    detail = excinfo.value.detail
    assert detail["reason"] == "pdf_parse_failed"
    assert detail["bytes"] == len(payload)
    error = detail["error"]
    assert isinstance(error, str)
    assert _CLASS_NAME_RE.match(error), f"{error!r} is a message, not a class name"
    rendered = f"{excinfo.value} {excinfo.value!r} {detail}"
    assert canary not in rendered


async def test_a_pdf_that_is_only_magic_bytes_still_fails_cleanly() -> None:
    """Detection routes on magic bytes, so five bytes are enough to reach the parser."""
    with pytest.raises(PipelineError) as excinfo:
        await extract_text("stub.pdf", b"%PDF-")

    assert excinfo.value.detail["reason"] == "pdf_parse_failed"


# ---------------------------------------------------------------------------
# 2. PDF: one page inside a readable document cannot be extracted
# ---------------------------------------------------------------------------


async def test_one_unreadable_page_does_not_lose_the_other_pages() -> None:
    """A page that raises during extraction costs that page only, not the document.

    Without the per-page guard the whole upload becomes ``pdf_parse_failed`` and
    two perfectly readable pages are thrown away with it.
    """
    raw = build_pdf_bytes(
        [["Page one synthetic marker"], ["Page two synthetic marker"], ["Page three marker"]],
        broken_pages={1},
    )

    text = await extract_text("damaged.pdf", raw)

    # The failed page contributes an empty part, so the two survivors keep their
    # order and stay separated rather than being run together.
    assert text == "Page one synthetic marker\n\n\nPage three marker\n"
    assert "Page two synthetic marker" not in text


async def test_a_pdf_whose_only_page_fails_is_empty_text_not_an_error() -> None:
    """No extractable text is a document with no text, not a rejected upload.

    Same rule a scanned PDF gets: the caller receives ``""`` and can decide, and
    ``put`` still stores the original bytes.
    """
    raw = build_pdf_bytes([["Unreachable synthetic marker"]], broken_pages={0})

    assert await extract_text("scan.pdf", raw) == ""


async def test_a_failing_page_is_logged_by_class_name_only(
    caplog: pytest.LogCaptureFixture,
) -> None:
    """The page-failure warning may carry an index and a class name, never a parser message.

    pypdf's message for this document names the object it choked on; a log line
    is the one place document bytes could leave this module without anybody
    reading a return value.
    """
    raw = build_pdf_bytes([["Good synthetic marker"], ["Lost synthetic marker"]], broken_pages={1})

    with caplog.at_level(logging.WARNING, logger=EXTRACTION_LOGGER):
        text = await extract_text("damaged.pdf", raw)

    assert "Good synthetic marker" in text
    messages = warnings_from(caplog, EXTRACTION_LOGGER)
    assert len(messages) == 1, messages
    message = messages[0]
    assert "page=1" in message
    assert "TypeError" in message
    # The pypdf message for this failure reads "argument of type 'NumberObject'
    # is not iterable" -- none of it may appear.
    assert "NumberObject" not in message
    assert "not iterable" not in message


# ---------------------------------------------------------------------------
# 3. Deviation 32: documents arrive hard-wrapped
# ---------------------------------------------------------------------------


async def test_pdf_prose_arrives_hard_wrapped() -> None:
    """A sentence drawn across two rows comes back with a newline inside it.

    Deviation 32 rests on this, and three modules are written against it: the
    detector's sentence-tail windows, the restorer, and ``textrules``' flexible
    whitespace. A ``_normalise`` that re-flowed single newlines into spaces would
    look like an improvement and would falsify all three at once.
    """
    raw = build_pdf_bytes([["The quarterly summary wraps", "across two rendered rows."]])

    text = await extract_text("wrapped.pdf", raw)

    assert "wraps\nacross" in text, text
    assert "wraps across" not in text
    # A hard wrap is a single newline: it must not be promoted to a paragraph
    # boundary either, because the detector chunks on those.
    assert "wraps\n\n" not in text


async def test_docx_line_break_arrives_hard_wrapped() -> None:
    """A manual line break inside a DOCX paragraph survives as a newline, not a space."""
    document = new_docx()
    paragraph = document.add_paragraph("The quarterly summary wraps")
    paragraph.add_run().add_break()
    paragraph.add_run("across a manual line break.")
    buffer = io.BytesIO()
    document.save(buffer)

    text = await extract_text("wrapped.docx", buffer.getvalue())

    assert "wraps\nacross" in text, text
    assert "wraps across" not in text


# ---------------------------------------------------------------------------
# 4. DOCX: the ZIP sniff when the central directory is gone
# ---------------------------------------------------------------------------


def truncate_to_local_headers(raw: bytes) -> bytes:
    """The front 60% of a ZIP: local headers intact, central directory gone."""
    return raw[: int(len(raw) * 0.6)]


def test_a_truncated_docx_is_still_recognised_as_a_word_package() -> None:
    """With no readable directory, the member-name probe is what keeps the routing right.

    ZIP member names are stored uncompressed in the local file headers, so a
    substring probe still answers "is this WordprocessingML" for a truncated
    upload. Getting this wrong is not cosmetic: the fallback returning ``False``
    routes a Word package into the *text* decoder, and the user gets a screenful
    of mojibake ZIP bytes back instead of an error.
    """
    truncated = truncate_to_local_headers(build_simple_docx_bytes())

    # Premise, established with the stdlib rather than with the module under
    # test: the directory really is unreadable, and the member name really is
    # still in the bytes.
    with pytest.raises(zipfile.BadZipFile):
        zipfile.ZipFile(io.BytesIO(truncated)).namelist()
    assert b"word/document.xml" in truncated
    assert truncated.startswith(b"PK\x03\x04")

    assert detect_kind("report.docx", truncated) == KIND_DOCX
    # ...and the bytes decide, so a mislabelled name does not change the answer.
    assert detect_kind("report.txt", truncated) == KIND_DOCX


async def test_a_truncated_docx_fails_as_a_docx_not_as_mojibake() -> None:
    """The size guard cannot read the directory either, and must not block on that.

    ``_declared_uncompressed_size`` returns ``None`` here; the extractor has to
    carry on to python-docx, which is what turns the upload into a structured
    ``docx_parse_failed`` rather than a silent success.
    """
    truncated = truncate_to_local_headers(build_simple_docx_bytes())

    with pytest.raises(PipelineError) as excinfo:
        await extract_text("report.docx", truncated)

    detail = excinfo.value.detail
    assert detail["reason"] == "docx_parse_failed"
    error = detail["error"]
    assert isinstance(error, str)
    assert _CLASS_NAME_RE.match(error), f"{error!r} is a message, not a class name"
    assert detail["bytes"] == len(truncated)


def test_a_truncated_zip_that_is_not_a_word_package_is_refused() -> None:
    """The fallback must answer "no" as confidently as it answers "yes".

    A probe that gave up and returned ``True`` would hand every ZIP -- an .xlsx,
    a .jar -- to python-docx and turn a clear ``unsupported_document_type`` into
    a parse failure.
    """
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>" * 200)
    truncated = truncate_to_local_headers(buffer.getvalue())

    with pytest.raises(zipfile.BadZipFile):
        zipfile.ZipFile(io.BytesIO(truncated)).namelist()

    with pytest.raises(PipelineError) as excinfo:
        detect_kind("workbook.docx", truncated)

    assert excinfo.value.detail["reason"] == "unsupported_document_type"
    assert excinfo.value.detail["detected"] == "zip_not_docx"


async def test_a_docx_with_a_corrupted_member_fails_structurally() -> None:
    """A readable directory over an unreadable member is still ``docx_parse_failed``.

    Different path from the truncated case: the size guard succeeds and returns a
    number, so the failure has to come out of the parse rather than the sniff.
    """
    raw = bytearray(build_simple_docx_bytes())
    middle = len(raw) // 2
    for offset in range(middle, middle + 64):
        raw[offset] ^= 0xFF
    corrupted = bytes(raw)

    # Premise: the directory is intact, so this is not the truncated case again.
    assert zipfile.ZipFile(io.BytesIO(corrupted)).namelist()

    with pytest.raises(PipelineError) as excinfo:
        await extract_text("report.docx", corrupted)

    assert excinfo.value.detail["reason"] == "docx_parse_failed"


# ---------------------------------------------------------------------------
# 5. TXT: the last rung of the decode ladder
# ---------------------------------------------------------------------------


async def test_a_bom_over_non_utf8_bytes_is_dropped_not_decoded() -> None:
    """A BOM in front of latin-1 text must not survive into the first token.

    ``utf-8-sig`` fails on the body, so the BOM has to be stripped by hand before
    the fallback runs. Leaving it attached decodes it as ``ï»¿`` and glues three
    characters onto the first word, which then shifts every detector offset in
    the document.
    """
    body = "Café rôle déjà vu."
    raw = b"\xef\xbb\xbf" + body.encode("latin-1")

    with pytest.raises(UnicodeDecodeError):
        raw.decode("utf-8-sig")

    text = await extract_text("legacy-bom.txt", raw)

    assert text == body + "\n"
    assert "\ufeff" not in text
    assert "ï»¿" not in text


async def test_undecodable_bytes_fall_back_to_latin1_not_to_replacement() -> None:
    """The final rung maps every byte to a character; it does not blank them out.

    ``utf-8`` with ``errors="replace"`` would also "not raise" while turning the
    high bytes into U+FFFD -- the same shape of output, with the content gone.
    """
    raw = b"header \xff\xfe tail"

    text = await extract_text("binaryish.txt", raw)

    assert text == "header ÿþ tail\n"
    assert "\ufffd" not in text


@pytest.mark.parametrize("filename", ["empty.txt", "empty.md", "empty.pdf", "empty.docx"])
async def test_a_zero_byte_upload_is_empty_text_for_every_supported_suffix(
    filename: str,
) -> None:
    """Zero bytes carry no magic, so every supported suffix decodes to ``""``.

    An empty upload is a document with no text, not a rejected one -- raising
    here would turn a harmless upload into a 4xx.
    """
    assert await extract_text(filename, b"") == ""


async def test_a_zero_byte_upload_with_an_unsupported_suffix_is_still_refused() -> None:
    """Emptiness does not excuse the type check: ``.exe`` is refused at zero bytes too."""
    with pytest.raises(PipelineError) as excinfo:
        await extract_text("payload.exe", b"")

    assert excinfo.value.detail["reason"] == "unsupported_document_type"
    assert excinfo.value.detail["suffix"] == ".exe"


async def test_an_unsupported_suffix_beats_perfectly_readable_text() -> None:
    """Readable UTF-8 under an unsupported name is refused, and the text stays out of the error."""
    canary = "SYNTHETIC-CANARY-TXT-8823"

    with pytest.raises(PipelineError) as excinfo:
        await extract_text("notes.exe", f"plain readable prose {canary}\n".encode())

    detail = excinfo.value.detail
    assert detail["reason"] == "unsupported_document_type"
    assert detail["detected"] == "unknown"
    assert canary not in f"{excinfo.value} {excinfo.value!r} {detail}"


def test_a_windows_path_with_a_dotted_directory_has_no_suffix() -> None:
    """Backslash separators are normalised before the suffix is taken.

    ``C:\\my.docs\\report`` has no extension. Reading it as a POSIX path without
    normalising gives the whole tail after the last dot -- ``.docs\\report`` --
    which is not a suffix, would never match ``SUPPORTED_SUFFIXES``, and would be
    echoed back inside a structured error detail.
    """
    with pytest.raises(PipelineError) as excinfo:
        detect_kind("C:\\my.docs\\report", b"plain readable prose\n")

    assert excinfo.value.detail["suffix"] == ""


# ---------------------------------------------------------------------------
# 6. document_store: the branches that only run when the filesystem misbehaves
# ---------------------------------------------------------------------------


@pytest.fixture
async def store(settings: Settings) -> Any:
    async with EncryptedDocumentStore(settings) as opened:
        yield opened


async def test_a_cleanup_failure_does_not_replace_the_upload_error(
    store: EncryptedDocumentStore,
    settings: Settings,
    monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture,
) -> None:
    """An undeletable orphan is logged, not raised: the caller keeps the real cause.

    ``put`` unwinds its own blob writes, and that unwind runs on the way out of a
    failure. If the unlink's own ``OSError`` escaped, a read-only or full data
    directory would replace the genuine write error with a confusing
    ``PermissionError`` from the cleanup path -- the diagnosis for a disk problem
    would point at the wrong syscall.
    """
    writes = 0
    real_write = store._write_blob

    async def fail_on_the_second_write(path: Path, blob: bytes) -> None:
        nonlocal writes
        writes += 1
        if writes == 2:
            raise OSError(5, "I/O error writing the text blob")
        await real_write(path, blob)

    async def unlink_is_refused(path: Any) -> None:
        raise PermissionError(13, "read-only file system")

    monkeypatch.setattr(store, "_write_blob", fail_on_the_second_write)
    monkeypatch.setattr(document_store.aiofiles.os, "remove", unlink_is_refused)

    with caplog.at_level(logging.WARNING, logger=STORE_LOGGER):
        with pytest.raises(OSError, match="I/O error writing the text blob"):
            await store.put("user-alpha-0001", "note.txt", b"synthetic body\n", "medical")

    messages = warnings_from(caplog, STORE_LOGGER)
    assert len(messages) == 2, messages
    kinds = sorted(message.rsplit("kind=", 1)[1].rstrip(")") for message in messages)
    assert kinds == sorted([RAW_SUFFIX, TEXT_SUFFIX])

    # The warning names the file that was left behind, so the operator can find
    # it -- and it really is still there, because the unlink was refused.
    doc_ids = {message.split("doc_id=", 1)[1].split(",", 1)[0] for message in messages}
    assert len(doc_ids) == 1
    doc_id = doc_ids.pop()
    assert DOC_ID_RE.match(doc_id), doc_id
    blob_dir = Path(settings.data_dir) / "documents"
    assert [path.name for path in blob_dir.iterdir()] == [f"{doc_id}{RAW_SUFFIX}"]


async def test_sweep_skips_a_blob_that_disappears_between_listing_and_stat(
    store: EncryptedDocumentStore, settings: Settings, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A file that vanishes mid-sweep is skipped; the rest of the sweep still runs.

    The directory listing and the age check are two syscalls apart, so a
    concurrent delete -- a second process, or an operator -- lands in between.
    Letting that ``OSError`` out would abort startup recovery partway through and
    leave every orphan after the vanished one on disk forever.
    """
    blob_dir = Path(settings.data_dir) / "documents"
    vanishing = blob_dir / f"{'a' * 32}{RAW_SUFFIX}"
    survivor = blob_dir / f"{'b' * 32}{RAW_SUFFIX}"
    for path in (vanishing, survivor):
        path.write_bytes(b"orphaned ciphertext stand-in\n")

    real_stat = document_store.aiofiles.os.stat

    async def stat_racing_a_delete(path: Any, *args: Any, **kwargs: Any) -> Any:
        if Path(path).name == vanishing.name:
            raise FileNotFoundError(2, "no such file")
        return await real_stat(path, *args, **kwargs)

    monkeypatch.setattr(document_store.aiofiles.os, "stat", stat_racing_a_delete)

    removed = await store.sweep_orphans(min_age_seconds=0)

    assert removed == 1
    assert [path.name for path in blob_dir.iterdir()] == [vanishing.name]
